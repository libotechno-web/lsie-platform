import io
import urllib.parse
from datetime import datetime
from flask import Response
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def create_excel_response(sheet_title, title_text, headers, rows, ascii_filename, arabic_filename=None):
    """
    توليد ملف Microsoft Excel (.xlsx) احترافي متوافق كلياً مع اللغة العربية (RTL)
    مع تنسيقات وتصميم متناسق مع هوية الجمعية الليبية للابتكار وريادة الأعمال.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = str(sheet_title)[:31]  # الحد الأقصى لاسم الورقة في إكسيل 31 حرفاً
    ws.views.sheetView[0].rightToLeft = True  # اتجاه من اليمين إلى اليسار
    
    font_family = "Arial"
    
    # خط الشعار والعنوان
    title_font = Font(name=font_family, size=14, bold=True, color="FFFFFF")
    title_fill = PatternFill(start_color="0D6E3F", end_color="0D6E3F", fill_type="solid")  # أخضر الجمعية
    
    # خط التاريخ والبيانات التوضيحية
    meta_font = Font(name=font_family, size=10, bold=False, color="475569")
    meta_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    
    # رؤوس الأعمدة
    header_font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="133B66", end_color="133B66", fill_type="solid")  # كحلي مؤسسي
    
    # خط البيانات
    data_font = Font(name=font_family, size=10, color="1E293B")
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    
    # حدود الخلايا
    thin_side = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(top=thin_side, left=thin_side, right=thin_side, bottom=thin_side)
    
    # محاذاة
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    right_align = Alignment(horizontal="right", vertical="center", wrap_text=True)
    
    num_cols = len(headers)
    last_col_letter = get_column_letter(num_cols)
    
    # 1. شريط العنوان الرئيسي
    ws.merge_cells(f"A1:{last_col_letter}1")
    cell_a1 = ws["A1"]
    cell_a1.value = f"الجمعية الليبية للابتكار وريادة الأعمال | {title_text}"
    cell_a1.font = title_font
    cell_a1.fill = title_fill
    cell_a1.alignment = center_align
    ws.row_dimensions[1].height = 40
    
    # 2. شريط التاريخ والتوثيق
    ws.merge_cells(f"A2:{last_col_letter}2")
    cell_a2 = ws["A2"]
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M")
    cell_a2.value = f"تاريخ التصدير: {now_str} | إجمالي السجلات: {len(rows)} | المنصة الوطنية للابتكار وريادة الأعمال"
    cell_a2.font = meta_font
    cell_a2.fill = meta_fill
    cell_a2.alignment = center_align
    ws.row_dimensions[2].height = 24
    
    # 3. صف فارغ فاصل
    ws.row_dimensions[3].height = 10
    
    # 4. رؤوس الأعمدة (الصف 4)
    ws.row_dimensions[4].height = 28
    for col_idx, h in enumerate(headers, 1):
        c = ws.cell(row=4, column=col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center_align
        c.border = cell_border
        
    # 5. صفوف البيانات (بدءاً من الصف 5)
    for row_idx, row_data in enumerate(rows, 5):
        ws.row_dimensions[row_idx].height = 22
        fill_to_use = zebra_fill if (row_idx % 2 == 0) else white_fill
        for col_idx, val in enumerate(row_data, 1):
            display_val = val if val is not None else "-"
            c = ws.cell(row=row_idx, column=col_idx, value=display_val)
            c.font = data_font
            c.fill = fill_to_use
            c.border = cell_border
            
            header_name = headers[col_idx - 1]
            if isinstance(val, (int, float)) or any(k in header_name for k in ['المعرف', '#', 'الهاتف', 'تاريخ', 'الحالة', 'النوع']):
                c.alignment = center_align
            else:
                c.alignment = right_align
                
    # 6. ضبط أبعاد الأعمدة تلقائياً بناءً على المحتوى
    for col in ws.columns:
        col_idx = col[0].column
        col_letter = get_column_letter(col_idx)
        # حساب أطول نص في العمود (من الصف 4 فما بعد لتجنب العنوان المدمج)
        max_len = 0
        for cell in col[3:]:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        adjusted_width = max(max_len + 5, 14)
        if adjusted_width > 50:
            adjusted_width = 50
        ws.column_dimensions[col_letter].width = adjusted_width
        
    # حفظ في ذاكرة وسيطة
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    # تأمين اسم الملف لترويسة HTTP بدون أي محارف غير لاتينية
    safe_ascii = ascii_filename if ascii_filename.endswith('.xlsx') else f"{ascii_filename}.xlsx"
    if not arabic_filename:
        arabic_filename = safe_ascii
    elif not arabic_filename.endswith('.xlsx'):
        arabic_filename = f"{arabic_filename}.xlsx"
        
    encoded_ar = urllib.parse.quote(arabic_filename)
    
    response = Response(
        output.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    # ترويسة آمنة تماماً: filename لاتيني نقي لمنع خطأ latin-1 codec، و filename* لاسم عربي في المتصفحات
    response.headers["Content-Disposition"] = f'attachment; filename="{safe_ascii}"; filename*=UTF-8\'\'{encoded_ar}'
    response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate"
    response.headers["Pragma"] = "no-cache"
    response.headers["Expires"] = "0"
    return response
